import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\PO-TC1ED\Downloads\20260505_ESS062_CANCER (5)_ListaErrores.xlsx')
for s in wb.sheetnames:
    print('===', s)
    ws = wb[s]
    for i,r in enumerate(ws.iter_rows(values_only=True)):
        print(r)
        if i>=60: break
    print()
